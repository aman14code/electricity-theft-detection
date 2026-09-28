import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
MD_FILE = os.path.join(SCRIPT_DIR, "comprehensive_study_guide.md")
OUT_DOCX = os.path.join(SCRIPT_DIR, "PowerGuard_Study_Guide.docx")

def create_word_doc():
    doc = Document()
    
    # Title
    title = doc.add_heading("PowerGuard — Complete Deep Dive Study Guide", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph("Electricity Theft Detection System", style="Subtitle")
    doc.add_page_break()

    with open(MD_FILE, "r", encoding="utf-8") as f:
        lines = f.readlines()

    in_code = False
    for line in lines:
        line = line.strip()
        
        # Code blocks
        if line.startswith("```"):
            in_code = not in_code
            continue
            
        if in_code:
            p = doc.add_paragraph(line)
            if p.runs:
                p.runs[0].font.name = 'Courier New'
                p.runs[0].font.size = Pt(9)
            continue
            
        # Headings
        if line.startswith("# ") and not line.startswith("## "):
            doc.add_heading(line.lstrip("# ").strip(), level=1)
        elif line.startswith("## "):
            doc.add_heading(line.lstrip("## ").strip(), level=2)
        elif line.startswith("### "):
            doc.add_heading(line.lstrip("### ").strip(), level=3)
        elif line.startswith("#### "):
            doc.add_heading(line.lstrip("#### ").strip(), level=4)
        
        # Bullets
        elif line.startswith("- ") or line.startswith("* "):
            doc.add_paragraph(line.lstrip("-* ").strip(), style='List Bullet')
            
        # Blockquotes
        elif line.startswith(">"):
            p = doc.add_paragraph(line.lstrip("> ").strip(), style='Quote')
            
        # Tables (Very basic handling for display)
        elif "|" in line and line.startswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            if all(c.replace("-", "").strip() == "" for c in cells):
                continue # Skip separator
            p = doc.add_paragraph(" | ".join(cells))
            p.runs[0].font.name = 'Consolas'
            p.runs[0].font.size = Pt(10)
            
        # Empty lines
        elif not line:
            pass
            
        # Normal text
        else:
            doc.add_paragraph(line)

    doc.save(OUT_DOCX)
    print(f"Created Word Document at {OUT_DOCX}")

if __name__ == "__main__":
    create_word_doc()
