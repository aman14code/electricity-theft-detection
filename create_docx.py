import docx
from docx.shared import Inches
import os

doc = docx.Document()

# Add the Prompt Text
text = """Please rewrite my draft research paper into a final, highly technical IEEE-format publication. Use the attached sample paper (Kawoosa et al.) as a structural guide, but strictly use MY project data, methodology, and results provided below.

CRITICAL FORMATTING INSTRUCTIONS:
1. IEEE Format: Use strict IEEE structure (two-column layout assumption, Roman numerals for main headings: I. INTRODUCTION, II. LITERATURE REVIEW, III. METHODOLOGY, IV. RESULTS, V. CONCLUSION).
2. Formulas: You MUST include LaTeX mathematical derivations for SMOTE, Soft Voting Ensemble, Accuracy, Precision, Recall, F1-Score, and AUC-ROC in their respective sections.
3. Figure Placeholders: Throughout the text, you MUST insert explicit placeholders (e.g., [INSERT FIG 1: CLASS IMBALANCE PIE CHART HERE]) exactly where the 10 images should go, accompanied by a formal IEEE-style figure caption.

### MY PROJECT FACT SHEET & FIGURE MAPPING

I. INTRODUCTION & II. LITERATURE REVIEW
* Context: SGCC dataset, severe NTL losses.
* Contribution: A 5-model Soft Voting Ensemble paired with a deployable Microservice architecture (React, Node.js, FastAPI) and a custom 8-measure heuristic explainability engine.

III. METHODOLOGY
* A. Dataset: 42,372 consumers, 1,034 days, 8.53% theft rate. -> [INSERT FIG 1 HERE]
* B. Data Pre-processing & Imbalance: Add mathematical derivation for SMOTE oversampling here. Mention Cost-Sensitive learning and PR-curve threshold optimization. We did NOT use artificial attacks. -> [INSERT FIG 2 HERE]
* C. Feature Engineering: 47 features (Statistical, Temporal, Anomaly, Comparative). -> [INSERT FIG 3 HERE]
* D. Model Paradigms: Briefly define RF, XGBoost, DNN, Isolation Forest, and LSTM mathematically.
* E. Ensemble Strategy (Core Contribution): Add the mathematical derivation for Soft Voting here.
* F. Explainability: Mention SHAP and the custom 8-measure heuristic engine. -> [INSERT FIG 4 HERE]

IV. RESULTS AND DISCUSSION
* Provide the LaTeX formulas for Precision, Recall, F1, and AUC.
* A. Model Performance: Present the exact results table: Random Forest (AUC 0.793), XGBoost (AUC 0.785), DNN (AUC 0.786), IsoForest (AUC 0.587), LSTM (AUC 0.746). 
* The Winner: Soft Voting Ensemble achieved 89.88% Accuracy, 0.411 Precision, 0.435 Recall, 0.423 F1-Score, and 0.809 AUC-ROC. -> [INSERT FIG 6 HERE] & [INSERT FIG 7 HERE]
* B. Cross-Validation Stability: 5-Fold Stratified. RF Mean Acc: 89.13% (+/-1.3%). XGBoost Mean Acc: 88.95% (+/-0.78%). -> [INSERT FIG 5 HERE]
* C. ROC and PR Curves: Discuss why PR is critical for imbalanced data. -> [INSERT FIG 8] & [INSERT FIG 9]
* D. Confusion Matrix Analysis: For the Ensemble: 7,302 TN, 450 FP, 408 FN, 315 TP. Discuss the trade-off of false positives vs field inspection costs. -> [INSERT FIG 10 HERE]

V. CONCLUSION
* Write in the past tense. Summarize the success of the 0.809 AUC-ROC ensemble and the successful deployment of the microservice dashboard architecture.
"""

doc.add_paragraph(text)
doc.add_page_break()

# Add Images
images = [
    ('fig1_class_imbalance.png', 'Fig 1: Class Imbalance'),
    ('fig2_smote_impact.png', 'Fig 2: SMOTE Impact'),
    ('fig3_consumption_profiles.png', 'Fig 3: Consumption Profiles'),
    ('fig4_feature_importance.png', 'Fig 4: Feature Importance'),
    ('fig5_cross_validation.png', 'Fig 5: Cross Validation Stability'),
    ('model_comparison_bar.png', 'Fig 6: Model Comparison Bar Chart'),
    ('fig7_radar_chart.png', 'Fig 7: Multidimensional Radar Chart'),
    ('roc_curves.png', 'Fig 8: ROC Curves'),
    ('fig6_pr_curves.png', 'Fig 9: Precision-Recall Curves'),
    ('confusion_matrix.png', 'Fig 10: Confusion Matrix')
]

graph_dir = "paper_graphs"
for img_name, title in images:
    img_path = os.path.join(graph_dir, img_name)
    if os.path.exists(img_path):
        doc.add_heading(title, level=2)
        doc.add_picture(img_path, width=Inches(6.0))
        doc.add_paragraph()

doc.save("Claude_Master_Prompt_With_Images.docx")
print("DOCX created successfully!")
